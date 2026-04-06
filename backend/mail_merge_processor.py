"""
Mail Merge Processor - Uses SmartMailMergeConverter for Vietnamese forms
"""
import uuid
from pathlib import Path
from typing import Dict
from smart_mail_merge_converter import SmartMailMergeConverter


class MailMergeProcessor:
    """Convert .docx to Mail Merge template using SmartMailMergeConverter"""

    def __init__(self, gemini_api_key: str = None, timeout: int = 30):
        """Initialize processor

        Args:
            gemini_api_key: Not used, kept for backward compatibility
            timeout: Not used, kept for backward compatibility
        """
        self.w_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def convert_to_mail_merge(self, docx_path: str, output_path: str = None) -> Dict:
        """Convert .docx to Mail Merge template using SmartMailMergeConverter

        Args:
            docx_path: Path to input .docx file
            output_path: Path to save template (optional)

        Returns:
            Dict with template_id, fields, and HTML preview
        """
        # Generate output path if not provided
        if output_path is None:
            template_id = str(uuid.uuid4())
            base_dir = Path(__file__).parent
            output_dir = base_dir / "uploads" / "templates"
            output_dir.mkdir(parents=True, exist_ok=True)
            output_path = output_dir / f"{template_id}.docx"
        else:
            output_path = Path(output_path)
            template_id = output_path.stem

        try:
            # Use SmartMailMergeConverter
            converter = SmartMailMergeConverter(docx_path)
            fields = converter.convert(str(output_path))

            # Generate HTML preview from converted template
            html_preview = self._generate_html_preview(str(output_path), fields)

            return {
                "template_id": template_id,
                "fields": fields,
                "field_count": len(fields),
                "html_preview": html_preview,
                "method": "smart_converter",
                "note": "Full XML surgical injection with Vietnamese support"
            }

        except Exception as e:
            raise RuntimeError(f"Conversion failed: {str(e)}")

    def _generate_html_preview(self, docx_path: str, fields: list) -> str:
        """Generate HTML preview with highlighted placeholders and preserved document order"""
        from docx import Document
        from docx.oxml.text.paragraph import CT_P
        from docx.oxml.table import CT_Tbl
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        try:
            doc = Document(docx_path)
            html_parts = []
            
            # Use body children directly to preserve document order
            for child in doc.element.body.iterchildren():
                if isinstance(child, CT_P):
                    para = Paragraph(child, doc)
                    html_parts.append(self._process_para_to_html(para))
                elif isinstance(child, CT_Tbl):
                    table = Table(child, doc)
                    html_parts.append(self._process_table_to_html(table))
            
            return "\n".join([p for p in html_parts if p])
            
        except Exception as e:
            import traceback
            return f'''
            <div class="error" style="padding: 20px; background-color: #fee; border: 1px solid #fcc; border-radius: 4px; color: #c33;">
                <h3>Lỗi chuyển đổi DOCX</h3>
                <p>Không thể chuyển đổi tài liệu sang HTML.</p>
                <pre style="background: #fff; padding: 10px; border-radius: 4px; overflow: auto;">{str(e)}\n{traceback.format_exc()}</pre>
            </div>
            '''

    def _get_run_style_text(self, rPr) -> str:
        """Map Word XML styles to CSS properties"""
        if rPr is None:
            return ""
            
        styles = []
        style_map = {
            f"{self.w_ns}b": "font-weight: bold",
            f"{self.w_ns}i": "font-style: italic",
            f"{self.w_ns}u": "text-decoration: underline",
            f"{self.w_ns}strike": "text-decoration: line-through"
        }
        
        for root_tag, css in style_map.items():
            if rPr.find(root_tag) is not None:
                styles.append(css)
                
        sz = rPr.find(f"{self.w_ns}sz")
        if sz is not None and sz.get(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}val"):
            styles.append(f"font-size: {int(sz.get(f'{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}val')) / 2}pt")
            
        color = rPr.find(f"{self.w_ns}color")
        if color is not None:
            val = color.get(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}val")
            if val and val != "auto" and len(val) == 6:
                styles.append(f"color: #{val}")
                
        shd = rPr.find(f"{self.w_ns}shd")
        if shd is not None:
            fill = shd.get(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}fill")
            if fill and fill != "auto":
                styles.append(f"background-color: #{fill}")
                
        rFonts = rPr.find(f"{self.w_ns}rFonts")
        if rFonts is not None:
            ascii_font = rFonts.get(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}ascii")
            if ascii_font:
                styles.append(f"font-family: '{ascii_font}', Calibri, Arial, sans-serif")
                
        return "; ".join(styles)

    def _process_xml_element_to_html(self, element) -> str:
        """Thực hiện xử lý XML element để lấy html"""
        import re
        if element is None:
            return ""
            
        tag_name = element.tag.split('}')[1] if '}' in element.tag else element.tag
        
        if tag_name == 'r':
            text_parts = []
            rPr = element.find(f"{self.w_ns}rPr")
            style_text = self._get_run_style_text(rPr)
            
            for child in element:
                c_tag = child.tag.split('}')[1] if '}' in child.tag else child.tag
                if c_tag == 't' and child.text:
                    text_parts.append(child.text)
                elif c_tag == 'tab':
                    text_parts.append("    ")
                elif c_tag == 'br':
                    text_parts.append("<br>")
                    
            text = "".join(text_parts)
            if not text:
                return ""
                
            def highlight_placeholder(match):
                field_name = match.group(1).strip()
                return f'<span class="mail-merge-placeholder" data-field="{field_name}" contenteditable="false" style="{style_text}">«{field_name}»</span>'
                
            text_with_highlights = re.sub(r'«([^»]+)»', highlight_placeholder, text)
            if style_text and "mail-merge-placeholder" not in text_with_highlights:
                return f'<span style="{style_text}">{text_with_highlights}</span>'
            return text_with_highlights
            
        elif tag_name == 'fldSimple':
            instr = element.get(f"{self.w_ns}instr", "")
            match = re.search(r'MERGEFIELD\s+(\S+)', instr)
            field_name = match.group(1) if match else "unknown"
            z_match = re.search(r'\\z\s*"([^"]*)"', instr)
            field_original = z_match.group(1) if z_match else ""
            
            nested_r = element.find(f"{self.w_ns}r")
            style_text = ""
            actual_text = ""
            if nested_r is not None:
                style_text = self._get_run_style_text(nested_r.find(f"{self.w_ns}rPr"))
                actual_text = "".join(t.text for t in nested_r.findall(f"{self.w_ns}t") if t.text)
                
            if actual_text and not re.match(r'^«[^»]+»$', actual_text):
                return f'<span style="{style_text}">{actual_text}</span>'
            return f'<span class="mail-merge-placeholder" data-field="{field_name}" data-original="{field_original}" contenteditable="false" style="{style_text}">«{field_name}»</span>'
            
        # Recursive process children
        return "".join(self._process_xml_element_to_html(child) for child in element)

    def _process_para_to_html(self, para) -> str:
        """Xử lý Paragraph thành thẻ p hoặc h1/h2/h3"""
        jc = para._p.find(f"{self.w_ns}pPr/{self.w_ns}jc")
        align = jc.get(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}val") if jc is not None else "left"
        align_map = {"center": "center", "right": "right", "both": "justify"}
        alignment = align_map.get(align, "left")
        
        style_name = para.style.name if para.style else "Normal"
        tag = "h1" if "Heading 1" in style_name else "h2" if "Heading 2" in style_name else "h3" if "Heading 3" in style_name else "p"
        
        content = "".join(self._process_xml_element_to_html(child) for child in para._p)
        if not content.strip():
            return "<p>&nbsp;</p>"
            
        align_style = f"text-align: {alignment};" if alignment != "left" else ""
        return f'<{tag} style="{align_style}">{content}</{tag}>'

    def _process_table_to_html(self, table) -> str:
        """Xử lý Table thành thẻ table html"""
        table_html = ['<table class="docx-table" style="border-collapse: collapse; width: 100%; margin: 10px 0;">']
        for row_idx, row in enumerate(table.rows):
            table_html.append('<tr style="border: 1px solid #ccc;">')
            for cell in row.cells:
                # Reuse process para for cells to support alignments inside table
                cell_content = "".join(self._process_para_to_html(para) for para in cell.paragraphs)
                tag = "th" if row_idx == 0 else "td"
                table_html.append(f'<{tag} style="border: 1px solid #ccc; padding: 5px;">{cell_content}</{tag}>')
            table_html.append('</tr>')
        table_html.append('</table>')
        return "\n".join(table_html)


